"""
excel_service.py
Exports application data to a multi-sheet Excel workbook using openpyxl.
"""
import os
from datetime import datetime, timezone
from openpyxl import Workbook
from openpyxl.styles import (Font, PatternFill, Alignment, Border, Side,
                              numbers)
from openpyxl.utils import get_column_letter
from app.models.application import Application
from app.models.business_details import BusinessDetails
from app.models.partner import Partner
from app.models.pharmacist import Pharmacist
from app.models.premises import Premises
from app.models.building_owner import BuildingOwner
from app.models.declaration import Declaration
from app.models.payment import Payment


NAVY = '1A3C6E'
GREEN = '28A745'
WHITE = 'FFFFFF'
LIGHT_GREY = 'F2F2F2'


def _header_style(cell):
    cell.font = Font(bold=True, color=WHITE, size=11)
    cell.fill = PatternFill('solid', fgColor=NAVY)
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)


def _alt_row(ws, row, col_count):
    fill = PatternFill('solid', fgColor=LIGHT_GREY)
    for c in range(1, col_count + 1):
        ws.cell(row=row, column=c).fill = fill


def _auto_width(ws, min_width=12, max_width=40):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = max(min_width, min(max_len + 2, max_width))


def _add_sheet_applications(wb, applications):
    ws = wb.create_sheet('Applications')
    headers = [
        'App Number', 'App Date', 'Licence Category', 'Licence Type',
        'Licence Forms', 'Constitution', 'Status', 'Created At', 'Updated At',
    ]
    ws.append(headers)
    for c, h in enumerate(headers, 1):
        _header_style(ws.cell(1, c))
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions

    for row_idx, app in enumerate(applications, start=2):
        ws.append([
            str(app.app_number),
            app.app_date.isoformat() if app.app_date else '',
            app.licence_category or '',
            app.licence_type or '',
            app.licence_forms or '',
            app.business_constitution or '',
            app.status or '',
            app.created_at.strftime('%Y-%m-%d %H:%M') if app.created_at else '',
            app.updated_at.strftime('%Y-%m-%d %H:%M') if app.updated_at else '',
        ])
        if row_idx % 2 == 0:
            _alt_row(ws, row_idx, len(headers))
    _auto_width(ws)


def _add_sheet_business(wb, applications):
    ws = wb.create_sheet('Business Details')
    headers = [
        'App Number', 'Firm Name', 'Trading Name', 'Constitution', 'Firm Reg No',
        'PAN', 'Email', 'Mobile', 'Alt Contact',
        'Door No', 'Floor', 'Street', 'Village', 'Gram Panchayat',
        'Mandal', 'District', 'State', 'PIN',
        'Proprietor Name', 'Proprietor Father Name',
        'Proprietor DOB', 'Proprietor Mobile', 'Proprietor PAN',
        'LLP Name', 'LLPIN', 'Company Name', 'CIN',
        'Deed Date', 'Deed Reg Details',
    ]
    ws.append(headers)
    for c, h in enumerate(headers, 1):
        _header_style(ws.cell(1, c))
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions

    for row_idx, app in enumerate(applications, start=2):
        bd = app.business_details
        if not bd:
            continue
        ws.append([
            str(app.app_number),
            bd.firm_name or '',
            bd.trading_name or '',
            app.business_constitution or '',
            bd.firm_reg_number or '',
            bd.pan or '',
            bd.email or '',
            bd.mobile or '',
            bd.alt_contact or '',
            bd.door_number or '',
            bd.floor or '',
            bd.street or '',
            bd.village or '',
            bd.gram_panchayat or '',
            bd.mandal or '',
            bd.district or '',
            bd.state or '',
            '\t' + (bd.pincode or ''),  # force text
            bd.proprietor_name or '',
            bd.proprietor_father_husband or '',
            bd.proprietor_dob.isoformat() if bd.proprietor_dob else '',
            bd.proprietor_mobile or '',
            bd.proprietor_pan or '',
            bd.llp_name or '',
            bd.llpin or '',
            bd.company_name or '',
            bd.cin or '',
            bd.deed_date.isoformat() if bd.deed_date else '',
            bd.deed_registration_details or '',
        ])
        if row_idx % 2 == 0:
            _alt_row(ws, row_idx, len(headers))
    _auto_width(ws)


def _add_sheet_partners(wb, applications):
    ws = wb.create_sheet('Partners')
    headers = [
        'App Number', 'Partner #', 'Full Name', 'Father Name',
        'Residential Address', 'Mobile', 'Email', 'PAN', 'Role',
        'Designated Partner', 'Authorized Signatory',
    ]
    ws.append(headers)
    for c, h in enumerate(headers, 1):
        _header_style(ws.cell(1, c))
    ws.freeze_panes = 'A2'

    row_idx = 2
    for app in applications:
        for p in app.partners:
            ws.append([
                str(app.app_number),
                p.partner_number,
                p.full_name or '',
                p.father_husband_name or '',
                p.residential_address or '',
                p.mobile or '',
                p.email or '',
                p.pan or '',
                p.role or '',
                'Yes' if p.is_designated_partner else 'No',
                'Yes' if p.is_authorized_signatory else 'No',
            ])
            if row_idx % 2 == 0:
                _alt_row(ws, row_idx, len(headers))
            row_idx += 1
    _auto_width(ws)


def _add_sheet_pharmacists(wb, applications):
    ws = wb.create_sheet('Pharmacists')
    headers = [
        'App Number', 'Full Name', 'Father Name', 'Qualification',
        'Reg Number', 'Reg Authority', 'Reg Validity',
        'Mobile', 'Email', 'Appointment Date', 'Employment Status',
    ]
    ws.append(headers)
    for c, h in enumerate(headers, 1):
        _header_style(ws.cell(1, c))
    ws.freeze_panes = 'A2'

    row_idx = 2
    for app in applications:
        for ph in app.pharmacists:
            ws.append([
                str(app.app_number),
                ph.full_name or '',
                ph.father_husband_name or '',
                ph.qualification or '',
                ph.registration_number or '',
                ph.registration_authority or '',
                ph.registration_validity.isoformat() if ph.registration_validity else '',
                ph.mobile or '',
                ph.email or '',
                ph.appointment_date.isoformat() if ph.appointment_date else '',
                ph.employment_status or '',
            ])
            if row_idx % 2 == 0:
                _alt_row(ws, row_idx, len(headers))
            row_idx += 1
    _auto_width(ws)


def _add_sheet_premises(wb, applications):
    ws = wb.create_sheet('Premises')
    headers = [
        'App Number', 'Ownership Type', 'Shop Name',
        'Door No', 'Floor', 'Street', 'Village', 'Mandal',
        'District', 'State', 'PIN', 'Area (sqm)', 'Storage Facilities',
    ]
    ws.append(headers)
    for c, h in enumerate(headers, 1):
        _header_style(ws.cell(1, c))
    ws.freeze_panes = 'A2'

    for row_idx, app in enumerate(applications, start=2):
        p = app.premises
        if not p:
            continue
        ws.append([
            str(app.app_number),
            p.ownership_type or '',
            p.shop_name or '',
            p.door_number or '',
            p.floor or '',
            p.street or '',
            p.village or '',
            p.mandal or '',
            p.district or '',
            p.state or '',
            '\t' + (p.pincode or ''),
            p.area_sqm or '',
            p.storage_facilities or '',
        ])
        if row_idx % 2 == 0:
            _alt_row(ws, row_idx, len(headers))
    _auto_width(ws)


def _add_sheet_building_owners(wb, applications):
    ws = wb.create_sheet('Building Owners')
    headers = [
        'App Number', 'Full Name', 'Father Name',
        'Residential Address', 'Mobile', 'Ownership Doc Details',
    ]
    ws.append(headers)
    for c, h in enumerate(headers, 1):
        _header_style(ws.cell(1, c))
    ws.freeze_panes = 'A2'

    row_idx = 2
    for app in applications:
        for bo in app.building_owners:
            ws.append([
                str(app.app_number),
                bo.full_name or '',
                bo.father_husband_name or '',
                bo.residential_address or '',
                bo.mobile or '',
                bo.ownership_doc_details or '',
            ])
            if row_idx % 2 == 0:
                _alt_row(ws, row_idx, len(headers))
            row_idx += 1
    _auto_width(ws)


def _add_sheet_declarations(wb, applications):
    ws = wb.create_sheet('Declarations')
    headers = ['App Number', 'Question Key', 'Question Text', 'Answer', 'Details']
    ws.append(headers)
    for c, h in enumerate(headers, 1):
        _header_style(ws.cell(1, c))
    ws.freeze_panes = 'A2'

    row_idx = 2
    for app in applications:
        for d in app.declarations:
            ws.append([
                str(app.app_number),
                d.question_key or '',
                d.question_text or '',
                d.answer or '',
                d.details or '',
            ])
            if row_idx % 2 == 0:
                _alt_row(ws, row_idx, len(headers))
            row_idx += 1
    _auto_width(ws)


def _add_sheet_payments(wb, applications):
    ws = wb.create_sheet('Payments')
    headers = [
        'App Number', 'Challan Number', 'Challan Date', 'Amount',
        'Payment Status', 'Payment Reference',
    ]
    ws.append(headers)
    for c, h in enumerate(headers, 1):
        _header_style(ws.cell(1, c))
    ws.freeze_panes = 'A2'

    row_idx = 2
    for app in applications:
        for pay in app.payments:
            ws.append([
                str(app.app_number),
                pay.challan_number or '',
                pay.challan_date.isoformat() if pay.challan_date else '',
                float(pay.amount) if pay.amount else '',
                pay.payment_status or '',
                pay.payment_reference or '',
            ])
            if row_idx % 2 == 0:
                _alt_row(ws, row_idx, len(headers))
            row_idx += 1
    _auto_width(ws)


def export_to_excel(applications: list, output_path: str) -> str:
    """
    Build an 8-sheet workbook from the supplied application list
    and save it to output_path.  Returns the path.
    """
    wb = Workbook()
    # Remove default sheet
    del wb[wb.sheetnames[0]]

    _add_sheet_applications(wb, applications)
    _add_sheet_business(wb, applications)
    _add_sheet_partners(wb, applications)
    _add_sheet_pharmacists(wb, applications)
    _add_sheet_premises(wb, applications)
    _add_sheet_building_owners(wb, applications)
    _add_sheet_declarations(wb, applications)
    _add_sheet_payments(wb, applications)

    wb.save(output_path)
    return output_path
